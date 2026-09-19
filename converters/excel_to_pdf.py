"""
Convert Excel spreadsheet (.xlsx, .xls) to PDF.
Supports LibreOffice and built-in Python vector table generator.
"""
import os
import shutil
import subprocess
import time
from typing import Dict, Any

try:
    from utils.system import find_libreoffice_binary
    from utils.fonts import resolve_font, init_system_fonts
except ImportError:
    from file_converters.utils.system import find_libreoffice_binary
    from file_converters.utils.fonts import resolve_font, init_system_fonts

def excel_to_pdf(input_xlsx: str, output_pdf: str) -> Dict[str, Any]:
    t0 = time.time()
    input_abs = os.path.abspath(input_xlsx)
    output_abs = os.path.abspath(output_pdf)

    if not os.path.isfile(input_abs):
        return {"success": False, "error": f"File not found: {input_xlsx}"}

    os.makedirs(os.path.dirname(output_abs) or ".", exist_ok=True)
    lo_bin = find_libreoffice_binary()

    # 1. Try LibreOffice if available
    if lo_bin:
        try:
            import tempfile
            with tempfile.TemporaryDirectory(prefix="excel2pdf_") as temp_dir:
                cmd = [lo_bin, "--headless", "--convert-to", "pdf", "--outdir", temp_dir, input_abs]
                res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=90)
                if res.returncode == 0:
                    base = os.path.splitext(os.path.basename(input_abs))[0]
                    cand = os.path.join(temp_dir, f"{base}.pdf")
                    if os.path.exists(cand):
                        shutil.move(cand, output_abs)
                        import pypdf
                        p_count = len(pypdf.PdfReader(output_abs).pages)
                        return {
                            "success": True,
                            "input": os.path.basename(input_xlsx),
                            "output": output_abs,
                            "count": f"{p_count} page(s)",
                            "size_bytes": os.path.getsize(output_abs),
                            "duration": time.time() - t0,
                            "engine": "libreoffice"
                        }
        except Exception:
            pass

    # 2. Pure Python Engine using openpyxl + reportlab
    try:
        import openpyxl
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
        from reportlab.lib.colors import Color, HexColor

        init_system_fonts()
        wb = openpyxl.load_workbook(input_abs, data_only=True)
        pdf_doc = SimpleDocTemplate(
            output_abs,
            pagesize=landscape(letter),
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        story = []
        page_width = 792 - 72  # 720 pt printable width

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if ws.max_row is None or ws.max_row == 0 or ws.max_column is None or ws.max_column == 0:
                continue

            title_style = ParagraphStyle(
                name=f"SheetTitle_{sheet_name}",
                fontName=resolve_font("Helvetica", is_bold=True),
                fontSize=14,
                leading=18,
                textColor=HexColor("#1E3A8A")
            )
            story.append(Paragraph(f"<b>Sheet: {sheet_name}</b>", title_style))
            story.append(Spacer(1, 10))

            table_data = []
            max_cols = min(ws.max_column, 25)
            max_rows = min(ws.max_row, 200)

            for r_idx in range(1, max_rows + 1):
                row_vals = []
                has_content = False
                for c_idx in range(1, max_cols + 1):
                    val = ws.cell(row=r_idx, column=c_idx).value
                    val_str = "" if val is None else str(val).strip()
                    if val_str:
                        has_content = True
                    if len(val_str) > 60:
                        val_str = val_str[:57] + "..."
                    val_str = val_str.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

                    cell_style = ParagraphStyle(
                        name=f"C_{sheet_name}_{r_idx}_{c_idx}",
                        fontName=resolve_font("Helvetica", is_bold=(r_idx == 1)),
                        fontSize=8.5 if r_idx == 1 else 8.0,
                        leading=10.5,
                        textColor=Color(0, 0, 0)
                    )
                    row_vals.append(Paragraph(f"<b>{val_str}</b>" if r_idx == 1 else val_str, cell_style))

                if has_content:
                    table_data.append(row_vals)

            if table_data:
                col_count = len(table_data[0])
                col_w = page_width / max(1, col_count)
                t = Table(table_data, colWidths=[col_w] * col_count)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), HexColor("#E0E7FF")),
                    ('INNERGRID', (0, 0), (-1, -1), 0.5, HexColor("#CBD5E1")),
                    ('BOX', (0, 0), (-1, -1), 1.0, HexColor("#64748B")),
                    ('TOPPADDING', (0, 0), (-1, -1), 4),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                    ('LEFTPADDING', (0, 0), (-1, -1), 4),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                ]))
                story.append(t)
                story.append(PageBreak())

        if not story:
            return {"success": False, "error": "Excel workbook is empty or has no readable data"}

        if isinstance(story[-1], PageBreak):
            story.pop()

        pdf_doc.build(story)
        import pypdf
        p_count = len(pypdf.PdfReader(output_abs).pages)

        return {
            "success": True,
            "input": os.path.basename(input_xlsx),
            "output": output_abs,
            "count": f"{p_count} page(s)",
            "size_bytes": os.path.getsize(output_abs),
            "duration": time.time() - t0,
            "engine": "native-python"
        }

    except Exception as e:
        return {"success": False, "error": f"Excel to PDF conversion error: {e}"}
