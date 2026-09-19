"""
Convert PDF document to PowerPoint presentation (.pptx).
Each PDF page becomes a high-resolution slide preserving exact visual fidelity.
"""
import io
import os
import time
from typing import Dict, Any
import pymupdf
from pptx import Presentation
from pptx.util import Pt

def pdf_to_ppt(input_pdf: str, output_pptx: str, dpi: int = 200) -> Dict[str, Any]:
    t0 = time.time()
    input_abs = os.path.abspath(input_pdf)
    output_abs = os.path.abspath(output_pptx)

    if not os.path.isfile(input_abs):
        return {"success": False, "error": f"PDF file not found: {input_pdf}"}

    os.makedirs(os.path.dirname(output_abs) or ".", exist_ok=True)

    try:
        doc = pymupdf.open(input_abs)
        page_count = len(doc)
        if page_count == 0:
            return {"success": False, "error": "PDF has no pages"}

        prs = Presentation()
        # Use dimensions from the first page
        first_page = doc[0]
        rect = first_page.rect
        prs.slide_width = Pt(rect.width)
        prs.slide_height = Pt(rect.height)

        blank_layout = prs.slide_layouts[6]

        for page in doc:
            slide = prs.slides.add_slide(blank_layout)
            pix = page.get_pixmap(dpi=dpi)
            img_bytes = pix.tobytes("png")

            img_stream = io.BytesIO(img_bytes)
            slide.shapes.add_picture(img_stream, 0, 0, width=Pt(page.rect.width), height=Pt(page.rect.height))

        doc.close()
        prs.save(output_abs)

        return {
            "success": True,
            "input": os.path.basename(input_pdf),
            "output": output_abs,
            "count": f"{page_count} slide(s)",
            "size_bytes": os.path.getsize(output_abs),
            "duration": time.time() - t0
        }
    except Exception as e:
        return {"success": False, "error": f"PDF to PPT conversion error: {e}"}
